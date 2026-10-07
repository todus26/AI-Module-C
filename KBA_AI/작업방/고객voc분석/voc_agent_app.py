"""고객사 VOC 분석 에이전트 (CrewAI + Gradio + Plotly)

실행:  python voc_agent_app.py

Agent 구성
  - VOC Agent    : CSV 로드, 통계 분석(Plotly 막대그래프), 워드클라우드
  - Issue Agent  : 산업군별 VOC 분석 -> 주요 이슈 도출 -> 개선과제 도출
  - Report Agent : VOC/Issue 분석 결과로 Word(.docx) 보고서 생성

사전 조건: 환경변수 OPENAI_API_KEY (선택: VOC_LLM_MODEL, 기본 gpt-4o-mini)
"""

import os
import re
import sys
import json
import time
import queue
import threading
import traceback
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

os.environ.setdefault("CREWAI_TRACING_ENABLED", "false")
os.environ.setdefault("CREWAI_DISABLE_TELEMETRY", "true")
os.environ.setdefault("OTEL_SDK_DISABLED", "true")

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import warnings

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402  Gradio 의 백엔드 복원과 워커 스레드 import 가 충돌하지 않도록 미리 로드

warnings.filterwarnings("ignore", message="function callbacks cannot be serialized")

import gradio as gr
import pandas as pd
import plotly.colors as px_colors
import plotly.graph_objects as go
from crewai import LLM, Agent, Crew, Process, Task
from crewai.tools import tool
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import BaseModel, Field
from wordcloud import WordCloud

# ════════════════════════════════════════════════════════════════════════════
# 1. 설정 / 한글 폰트
# ════════════════════════════════════════════════════════════════════════════
BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "outputs"
SAMPLE_CSV = BASE_DIR / "sample_voc.csv"
LLM_MODEL = os.getenv("VOC_LLM_MODEL", "gpt-4o-mini")

REQUIRED_COLUMNS = ["순번", "일자", "고객명", "산업군", "지역", "제품명", "분야", "불만"]
STAT_DIMS = ["산업군", "제품명", "분야"]

PLOTLY_FONT = "Malgun Gothic, 맑은 고딕, NanumGothic, Apple SD Gothic Neo, Noto Sans CJK KR, sans-serif"
DOCX_FONT = "맑은 고딕"


def find_korean_font() -> str | None:
    """워드클라우드/matplotlib 용 한글 폰트 파일 경로를 찾는다."""
    candidates = [
        r"C:\Windows\Fonts\malgun.ttf",
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/Library/Fonts/NanumGothic.ttf",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    from matplotlib import font_manager

    keys = ("malgun", "nanum", "gothic", "noto sans cjk kr", "apple sd")
    for f in font_manager.fontManager.ttflist:
        if any(k in f.name.lower() for k in keys):
            return f.fname
    return None


FONT_PATH = find_korean_font()


class UserError(Exception):
    """사용자에게 그대로 보여줄 오류."""


# ════════════════════════════════════════════════════════════════════════════
# 2. 세션 컨텍스트 / 로그
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class StatResult:
    fig: go.Figure
    table: pd.DataFrame
    text: str


@dataclass
class WcResult:
    path: str
    keywords: list[tuple[str, int]]
    industry: str
    text: str


@dataclass
class VocContext:
    df: pd.DataFrame
    filename: str
    workdir: Path
    stats: dict[str, StatResult] = field(default_factory=dict)
    wc: dict[str, WcResult] = field(default_factory=dict)
    wc_params: dict[str, Any] = field(default_factory=lambda: {"industry": "전체", "max_words": 80, "extra_stop": []})
    called: set[str] = field(default_factory=set)
    issues: "IssueAnalysis | None" = None
    report_content: "ReportContent | None" = None
    report_path: str | None = None


class JobLog:
    """백그라운드 작업의 진행률/로그를 UI 로 전달한다."""

    def __init__(self) -> None:
        self.lines: list[str] = []
        self.pct = 0
        self.status = "대기 중"
        self.failed = False
        self.version = 0
        self._lock = threading.Lock()

    def log(self, msg: str) -> None:
        with self._lock:
            self.lines.append(f"[{datetime.now():%H:%M:%S}] {msg}")
            self.version += 1

    def step(self, pct: int, msg: str) -> None:
        self.pct, self.status = pct, msg
        self.log(msg)

    def fail(self, msg: str) -> None:
        self.failed = True
        self.status = "오류 발생"
        self.log(f"[오류] {msg}")

    def text(self) -> str:
        with self._lock:
            return "\n".join(self.lines)


def progress_html(pct: int, status: str, failed: bool = False) -> str:
    color = "#e5484d" if failed else "#4f46e5"
    return (
        '<div style="margin:2px 0 6px 0;">'
        f'<div style="display:flex;justify-content:space-between;font-size:13px;margin-bottom:4px;">'
        f"<span><b>진행사항</b> · {status}</span><span>{pct}%</span></div>"
        '<div style="background:#e5e7eb;border-radius:8px;height:12px;overflow:hidden;">'
        f'<div style="width:{pct}%;height:100%;background:{color};transition:width .3s;"></div>'
        "</div></div>"
    )


def stream_job(work, final_fn, n_extra: int):
    """work(lg) 를 스레드로 실행하며 진행률/로그를 Gradio 로 스트리밍한다.

    출력 순서: [진행바, 로그, *메뉴별 결과(n_extra 개)]
    """
    lg = JobLog()
    holder: dict[str, Any] = {}

    def runner():
        try:
            holder["res"] = work(lg)
        except UserError as e:
            lg.fail(str(e))
        except Exception as e:  # noqa: BLE001
            lg.fail(f"{type(e).__name__}: {e}")
            lg.log(traceback.format_exc(limit=4))

    noop = [gr.update()] * n_extra
    thread = threading.Thread(target=runner, daemon=True)
    thread.start()

    def head():
        return [progress_html(lg.pct, lg.status, lg.failed), lg.text()]

    yield head() + noop
    last = -1
    while thread.is_alive():
        if lg.version != last:
            last = lg.version
            yield head() + noop
        time.sleep(0.3)
    thread.join()

    if lg.failed or "res" not in holder:
        yield head() + noop
        return
    lg.pct, lg.status = 100, "완료"
    lg.log("완료")
    yield head() + list(final_fn(holder["res"]))


# ════════════════════════════════════════════════════════════════════════════
# 3. 데이터 로드
# ════════════════════════════════════════════════════════════════════════════
def load_voc_csv(path: str, lg: JobLog) -> VocContext:
    lg.step(10, "[VOC Agent] CSV 파일 읽는 중")
    df = None
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            df = pd.read_csv(path, encoding=enc)
            lg.log(f"[VOC Agent] 인코딩 확인: {enc}")
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if df is None:
        raise UserError("CSV 인코딩을 인식할 수 없습니다. UTF-8 또는 CP949 로 저장해 주세요.")

    df.columns = [str(c).strip() for c in df.columns]
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise UserError(f"필수 컬럼이 없습니다: {', '.join(missing)} (필요: {', '.join(REQUIRED_COLUMNS)})")

    lg.step(50, "[VOC Agent] 데이터 정제 중")
    df = df[REQUIRED_COLUMNS].copy()
    for col in ["고객명", "산업군", "지역", "제품명", "분야", "불만"]:
        df[col] = df[col].fillna("").astype(str).str.strip().replace("", "미분류" if col != "불만" else "")
    df["일자"] = df["일자"].astype(str).str.strip()
    df = df.dropna(how="all").reset_index(drop=True)
    if df.empty:
        raise UserError("데이터가 비어 있습니다.")

    workdir = OUTPUT_DIR / f"{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}"
    workdir.mkdir(parents=True, exist_ok=True)
    lg.step(90, f"[VOC Agent] 로드 완료: {len(df):,}건")
    return VocContext(df=df, filename=Path(path).name, workdir=workdir)


def period_text(df: pd.DataFrame) -> str:
    dates = pd.to_datetime(df["일자"], errors="coerce").dropna()
    if dates.empty:
        return "-"
    return f"{dates.min():%Y-%m-%d} ~ {dates.max():%Y-%m-%d}"


def overview_md(ctx: VocContext) -> str:
    df = ctx.df
    return (
        f"**파일**: {ctx.filename} · **총 건수**: {len(df):,}건 · **기간**: {period_text(df)}  \n"
        f"**산업군** {df['산업군'].nunique()}개 · **제품** {df['제품명'].nunique()}개 · "
        f"**분야** {df['분야'].nunique()}개 · **고객사** {df['고객명'].nunique()}곳"
    )


# ════════════════════════════════════════════════════════════════════════════
# 4. VOC 통계 / 워드클라우드 (VOC Agent 의 핵심 기능)
# ════════════════════════════════════════════════════════════════════════════
def style_fig(fig: go.Figure, title: str, height: int = 430) -> go.Figure:
    fig.update_layout(
        title=dict(text=title, font=dict(size=17)),
        template="plotly_white",
        font=dict(family=PLOTLY_FONT, size=13),
        height=height,
        margin=dict(l=60, r=30, t=70, b=70),
        showlegend=False,
    )
    return fig


def compute_stats(ctx: VocContext, dim: str) -> str:
    counts = ctx.df[dim].value_counts()
    total = int(counts.sum())
    table = pd.DataFrame(
        {"구분": dim, "항목": counts.index.astype(str), "건수": counts.values, "비율(%)": counts.values / total * 100}
    )
    palette = px_colors.qualitative.Plotly
    fig = go.Figure(
        go.Bar(
            x=table["항목"],
            y=table["비율(%)"],
            text=[f"{v:.2f}%" for v in table["비율(%)"]],
            textposition="outside",
            customdata=table["건수"],
            hovertemplate="%{x}<br>비율 %{y:.2f}%<br>건수 %{customdata:,}건<extra></extra>",
            marker_color=[palette[i % len(palette)] for i in range(len(table))],
        )
    )
    style_fig(fig, f"{dim}별 VOC 비율 (총 {total:,}건)")
    fig.update_yaxes(title_text="비율(%)", ticksuffix="%", tickformat=".2f", range=[0, table["비율(%)"].max() * 1.2])
    fig.update_xaxes(title_text=dim, type="category", tickangle=-30 if len(table) > 6 else 0)

    items = ", ".join(f"{r.항목} {r._4:.2f}%({r.건수:,}건)" for r in table.head(10).itertuples())
    text = f"[{dim}별 VOC 비율] 총 {total:,}건 / 상위 항목: {items}"
    ctx.stats[dim] = StatResult(fig, table, text)
    return text


def format_stats_table(results: list[StatResult]) -> pd.DataFrame:
    frames = [r.table.assign(**{"비율(%)": r.table["비율(%)"].map(lambda v: f"{v:.2f}%"), "건수": r.table["건수"].map(lambda v: f"{v:,}")}) for r in results]
    return pd.concat(frames, ignore_index=True)


KO_STOPWORDS = set(
    "너무 정말 매우 계속 자꾸 항상 여러 대한 통해 관련 경우 때문 이후 현재 최근 또한 그리고 하지만 그러나 있는 없는 있어 없어 "
    "있음 없음 같은 같이 많이 아주 조금 다시 또는 이런 저희 우리 지속 발생 요청 문의 진행 확인 필요 불만 사용 제공 문제 부분 정도 "
    "대비 이상 이하 모든 해당 가장 매번 자주 빈번 하는 하고 해서 있습니다 없습니다 합니다 입니다 사항 내용 상황 대해 위해 따라 "
    "로서 에서 으로 에게 그냥 대부분 일부 전혀 다른 어려움 부담 불편 않아 않고 않는 않은 않았 못해 못한 지키 지키지".split()
)
KO_ENDINGS = sorted(
    "했습니다 합니다 됩니다 입니다 습니다 에서는 에서도 으로는 으로도 에서 으로 에게 까지 부터 보다 처럼 마다 이나 이며 이고 이다 "
    "하고 해서 하여 하는 하지 하게 되어 되는 되고 돼서 됨 함 에는 에도 에만 이라고 이라 라고 한다 된다 인데 지만 으며 거나 였다 했다 됐다 "
    "적인 적으로 한 된".split(),
    key=len,
    reverse=True,
)
KO_JOSA_STRONG = set("을를은는이가에")  # 명사 말음으로 드물어 항상 제거
KO_JOSA_WEAK = set("의도로과와만")  # 신뢰도/만족도 등 명사와 혼동되므로 어간이 단독 출현할 때만 제거
KO_PROTECTED = {"전문가", "사업가", "투자가", "손잡이", "신뢰도", "만족도", "정확도", "인지도"}
TOKEN_RE = re.compile(r"[가-힣]{2,}|[A-Za-z][A-Za-z0-9\-]+")


def extract_keywords(texts, extra_stop=()) -> Counter:
    """형태소 분석기 없이 동작하는 경량 한국어 키워드 추출.

    조사/어미를 제거하고(1글자 조사는 어간이 단독으로도 등장할 때만), 불용어를 제외한다.
    """
    texts = [str(t) for t in texts if str(t).strip()]
    raw = Counter(tok for t in texts for tok in TOKEN_RE.findall(t))
    stop = KO_STOPWORDS | {s.strip() for s in extra_stop if s.strip()}

    def normalize(tok: str) -> str | None:
        if not re.fullmatch(r"[가-힣]+", tok):
            return tok.upper() if len(tok) <= 4 else tok
        if tok.endswith(("요", "다", "까", "죠")):
            return None
        for end in KO_ENDINGS:
            if tok.endswith(end) and len(tok) - len(end) >= 2:
                return tok[: -len(end)]
        if len(tok) >= 3 and tok not in KO_PROTECTED:
            if tok[-1] in KO_JOSA_STRONG or (tok[-1] in KO_JOSA_WEAK and tok[:-1] in raw):
                return tok[:-1]
        return tok

    result: Counter = Counter()
    for tok, n in raw.items():
        word = normalize(tok)
        if word and len(word) >= 2 and word not in stop and tok not in stop:
            result[word] += n
    min_count = 2 if len(texts) >= 100 else 1
    return Counter({k: v for k, v in result.items() if v >= min_count})


def build_wordcloud(ctx: VocContext, industry: str = "전체", max_words: int = 80, extra_stop=()) -> str:
    if not FONT_PATH:
        raise UserError("한글 폰트를 찾을 수 없습니다. 맑은 고딕/나눔고딕 등을 설치해 주세요.")
    df = ctx.df if industry == "전체" else ctx.df[ctx.df["산업군"] == industry]
    freq = extract_keywords(df["불만"], extra_stop)
    if not freq:
        raise UserError("추출된 키워드가 없습니다. '불만' 컬럼 내용을 확인해 주세요.")
    top = freq.most_common(int(max_words))
    wc = WordCloud(
        font_path=FONT_PATH,
        width=1200,
        height=600,
        background_color="white",
        colormap="viridis",
        prefer_horizontal=0.95,
        max_font_size=130,
        relative_scaling=0.4,
        margin=4,
        max_words=int(max_words),
    ).generate_from_frequencies(dict(top))
    safe = re.sub(r"[^0-9A-Za-z가-힣_-]", "_", industry)
    path = ctx.workdir / f"wordcloud_{safe}.png"
    wc.to_file(str(path))
    brief = ", ".join(f"{w}({c})" for w, c in top[:20])
    text = f"[불만 키워드 - {industry}] 대상 {len(df):,}건 / 상위 키워드: {brief}"
    ctx.wc[industry] = WcResult(str(path), top, industry, text)
    return text


def keyword_bar(res: WcResult, n: int = 20) -> go.Figure:
    top = res.keywords[:n][::-1]
    fig = go.Figure(
        go.Bar(
            x=[c for _, c in top],
            y=[w for w, _ in top],
            orientation="h",
            text=[f"{c:,}" for _, c in top],
            textposition="outside",
            marker_color="#4f46e5",
            hovertemplate="%{y}: %{x:,}회<extra></extra>",
        )
    )
    style_fig(fig, f"상위 키워드 빈도 ({res.industry})", height=max(360, 22 * len(top) + 120))
    fig.update_xaxes(title_text="언급 횟수")
    return fig


def industry_summary_text(ctx: VocContext, samples: int = 6) -> str:
    """Issue Agent 용: 산업군별 VOC 요약."""
    df, total = ctx.df, len(ctx.df)
    blocks = []
    for ind, g in df.groupby("산업군"):
        blocks.append((len(g), ind, g))
    blocks.sort(key=lambda b: b[0], reverse=True)
    out = []
    for n, ind, g in blocks:
        fields = ", ".join(f"{k} {v / n * 100:.2f}%" for k, v in g["분야"].value_counts().head(3).items())
        prods = ", ".join(f"{k}({v}건)" for k, v in g["제품명"].value_counts().head(3).items())
        kws = ", ".join(f"{w}({c})" for w, c in extract_keywords(g["불만"]).most_common(8))
        examples = " / ".join(t[:90] for t in g["불만"].value_counts().head(samples).index)
        out.append(
            f"■ {ind}: {n:,}건 (전체의 {n / total * 100:.2f}%)\n"
            f"  - 주요 분야: {fields}\n  - 주요 제품: {prods}\n  - 핵심 키워드: {kws}\n  - 대표 불만: {examples}"
        )
    return "\n".join(out)


def stats_brief(ctx: VocContext, dim: str) -> str:
    if dim not in ctx.stats:
        compute_stats(ctx, dim)
    return ctx.stats[dim].text


def cross_brief(ctx: VocContext) -> str:
    lines = []
    for ind, g in ctx.df.groupby("산업군"):
        top = ", ".join(f"{k} {v / len(g) * 100:.2f}%" for k, v in g["분야"].value_counts().head(2).items())
        lines.append(f"{ind}: {top}")
    return "; ".join(lines)


# ════════════════════════════════════════════════════════════════════════════
# 5. CrewAI: 데이터 모델 / 도구 / Agent
# ════════════════════════════════════════════════════════════════════════════
class Improvement(BaseModel):
    priority: str = Field(description="우선순위: 상/중/하")
    action: str = Field(description="구체적인 개선과제")


class IndustryIssue(BaseModel):
    industry: str = Field(description="산업군명")
    voc_summary: str = Field(description="해당 산업군의 주요 VOC 요약 (1~2문장)")
    issues: list[str] = Field(default_factory=list, description="주요 이슈 2~3개")
    improvements: list[Improvement] = Field(default_factory=list, description="개선과제 2~3개")


class IssueAnalysis(BaseModel):
    overall_summary: str = Field(description="전체 VOC 이슈 총평 (2~3문장)")
    industries: list[IndustryIssue] = Field(default_factory=list)
    common_issues: list[str] = Field(default_factory=list, description="산업군 공통 이슈")


class ActionItem(BaseModel):
    priority: str = Field(description="우선순위: 상/중/하")
    industry: str = Field(description="대상 산업군 (공통이면 '공통')")
    issue: str = Field(description="대응 대상 이슈")
    action: str = Field(description="구체적인 대응방안")
    timeline: str = Field(description="추진 일정 (예: 즉시, 1개월, 분기 내)")
    effect: str = Field(description="기대효과")


class ReportContent(BaseModel):
    title: str = Field(description="보고서 제목")
    executive_summary: str = Field(description="핵심 요약. 줄바꿈으로 구분된 3~5개 불릿")
    industry_comment: str = Field(description="산업군별 VOC 통계 해설")
    field_comment: str = Field(description="분야별 VOC 통계 해설")
    wordcloud_comment: str = Field(description="불만 키워드(워드클라우드) 해설")
    action_plan: list[ActionItem] = Field(default_factory=list, description="대응방안 6~10개")
    conclusion: str = Field(description="결론 및 제언")


def get_llm() -> LLM:
    return LLM(model=LLM_MODEL, temperature=0.2)


def require_api_key() -> None:
    if not os.getenv("OPENAI_API_KEY"):
        raise UserError("환경변수 OPENAI_API_KEY 가 설정되어 있지 않습니다.")


def make_voc_tools(ctx: VocContext, lg: JobLog):
    @tool("analyze_voc_stats")
    def analyze_voc_stats(dimension: str) -> str:
        """VOC 데이터를 지정한 기준으로 집계해 건수와 비율(%)을 계산하고 막대그래프를 생성합니다.
        dimension 은 반드시 '산업군', '제품명', '분야' 중 하나여야 합니다."""
        dim = str(dimension).strip()
        if dim not in STAT_DIMS:
            return f"오류: dimension 은 {STAT_DIMS} 중 하나여야 합니다."
        lg.log(f"[VOC Agent] 도구 실행 - {dim}별 통계 계산 및 막대그래프 생성")
        ctx.called.add("analyze_voc_stats")
        return compute_stats(ctx, dim)

    @tool("generate_wordcloud")
    def generate_wordcloud() -> str:
        """VOC 의 '불만' 텍스트에서 키워드를 추출해 한글 워드클라우드 이미지를 생성하고 상위 키워드를 반환합니다."""
        p = ctx.wc_params
        lg.log(f"[VOC Agent] 도구 실행 - 키워드 추출 및 워드클라우드 생성 (산업군: {p['industry']})")
        ctx.called.add("generate_wordcloud")
        return build_wordcloud(ctx, p["industry"], p["max_words"], p["extra_stop"])

    return [analyze_voc_stats, generate_wordcloud]


def build_voc_agent(ctx: VocContext, lg: JobLog) -> Agent:
    return Agent(
        role="VOC 분석 전문가",
        goal="고객사 VOC 데이터를 통계와 키워드로 분석하고 핵심 인사이트를 정리한다",
        backstory="B2B 고객 불만 데이터를 다뤄 온 데이터 분석가. 반드시 제공된 도구로 수치를 산출하고, 수치는 소수점 둘째 자리까지 표기한다.",
        tools=make_voc_tools(ctx, lg),
        llm=get_llm(),
        verbose=False,
        allow_delegation=False,
        max_iter=8,
    )


def make_issue_tools(ctx: VocContext, lg: JobLog):
    @tool("summarize_voc_by_industry")
    def summarize_voc_by_industry() -> str:
        """산업군별 VOC 건수/비율, 주요 분야·제품, 핵심 키워드, 대표 불만 사례를 요약해 반환합니다."""
        lg.log("[Issue Agent] 도구 실행 - 산업군별 VOC 요약 집계")
        ctx.called.add("summarize_voc_by_industry")
        return industry_summary_text(ctx)

    return [summarize_voc_by_industry]


def parse_model(task: Task, model: type[BaseModel]):
    out = task.output
    if out is None:
        raise UserError(f"{model.__name__} 결과를 받지 못했습니다.")
    if getattr(out, "pydantic", None) is not None:
        return out.pydantic
    m = re.search(r"\{.*\}", out.raw or "", re.S)
    if not m:
        raise UserError(f"{model.__name__} 형식으로 변환할 수 없습니다.")
    return model.model_validate_json(m.group(0))


def no_braces(text: str) -> str:
    return text.replace("{", "(").replace("}", ")")


def run_crew(agent: Agent, tasks: list[Task]) -> Any:
    return Crew(agents=[agent], tasks=tasks, process=Process.sequential, verbose=False, tracing=False).kickoff()


# ── VOC Agent 실행 ───────────────────────────────────────────────────────────
def auto_insight(ctx: VocContext, dims: list[str]) -> str:
    lines = []
    for d in dims:
        if d in ctx.stats:
            t = ctx.stats[d].table.iloc[0]
            lines.append(f"- {d}: '{t['항목']}' 항목이 {t['비율(%)']:.2f}%({int(t['건수']):,}건)로 가장 높습니다.")
    return "\n".join(lines)


def run_voc_stats(ctx: VocContext, dims: list[str], lg: JobLog) -> str:
    for d in dims:
        ctx.stats.pop(d, None)
    insight = ""
    lg.step(15, "[VOC Agent] 에이전트 구성")
    if os.getenv("OPENAI_API_KEY"):
        try:
            agent = build_voc_agent(ctx, lg)
            task = Task(
                description=(
                    f"업로드된 VOC 데이터(총 {len(ctx.df):,}건)를 다음 기준별로 분석하세요: {', '.join(dims)}.\n"
                    "1) 각 기준마다 analyze_voc_stats 도구를 한 번씩 호출하세요. dimension 인자는 기준 이름을 정확히 전달합니다.\n"
                    "2) 도구 결과 수치만 사용하여 핵심 인사이트를 한국어 불릿 3~5개로 요약하세요. 비율은 소수점 둘째 자리까지 표기합니다."
                ),
                expected_output="기준별 핵심 수치와 인사이트 불릿 3~5개",
                agent=agent,
            )
            lg.step(30, "[VOC Agent] CrewAI 실행 중 (통계 도구 호출 + 인사이트 작성)")
            insight = str(run_crew(agent, [task]).raw)
        except Exception as e:  # noqa: BLE001
            lg.log(f"[경고] CrewAI 실행 실패, 직접 계산으로 대체합니다: {type(e).__name__}: {e}")
    else:
        lg.log("[경고] OPENAI_API_KEY 없음 - AI 해석 없이 직접 계산합니다.")

    lg.step(80, "[VOC Agent] 통계 결과 검증")
    for d in dims:
        if d not in ctx.stats:
            lg.log(f"[VOC Agent] '{d}' 통계를 직접 계산합니다.")
            compute_stats(ctx, d)
    return insight or auto_insight(ctx, dims)


def run_voc_wordcloud(ctx: VocContext, params: dict, lg: JobLog) -> tuple[WcResult, str]:
    ctx.wc_params = params
    ctx.wc.pop(params["industry"], None)
    insight = ""
    lg.step(15, "[VOC Agent] 에이전트 구성")
    if os.getenv("OPENAI_API_KEY"):
        try:
            agent = build_voc_agent(ctx, lg)
            task = Task(
                description=(
                    f"VOC '불만' 텍스트(대상 산업군: {params['industry']})의 키워드를 분석하세요.\n"
                    "1) generate_wordcloud 도구를 정확히 한 번 호출하세요.\n"
                    "2) 반환된 상위 키워드를 근거로 고객이 반복적으로 호소하는 주제를 한국어 불릿 3~4개로 해석하세요."
                ),
                expected_output="상위 키워드 기반 불만 주제 해석 불릿 3~4개",
                agent=agent,
            )
            lg.step(30, "[VOC Agent] CrewAI 실행 중 (키워드 추출 + 워드클라우드)")
            insight = str(run_crew(agent, [task]).raw)
        except Exception as e:  # noqa: BLE001
            lg.log(f"[경고] CrewAI 실행 실패, 직접 생성으로 대체합니다: {type(e).__name__}: {e}")
    else:
        lg.log("[경고] OPENAI_API_KEY 없음 - AI 해석 없이 직접 생성합니다.")

    lg.step(80, "[VOC Agent] 워드클라우드 결과 검증")
    if params["industry"] not in ctx.wc:
        lg.log("[VOC Agent] 워드클라우드를 직접 생성합니다.")
        build_wordcloud(ctx, params["industry"], params["max_words"], params["extra_stop"])
    res = ctx.wc[params["industry"]]
    if not insight:
        insight = "- 상위 키워드: " + ", ".join(f"{w}({c})" for w, c in res.keywords[:10])
    return res, insight


# ── Issue Agent 실행 ─────────────────────────────────────────────────────────
def run_issue_agent(ctx: VocContext, lg: JobLog, base_pct: int, span: int) -> IssueAnalysis:
    agent = Agent(
        role="VOC 이슈 분석가",
        goal="산업군별 VOC 에서 주요 이슈와 실행 가능한 개선과제를 도출한다",
        backstory="B2B 고객 경험 개선 컨설턴트. 데이터에 근거한 이슈만 제시하고, 개선과제는 구체적 행동 단위로 제안한다.",
        tools=make_issue_tools(ctx, lg),
        llm=get_llm(),
        verbose=False,
        allow_delegation=False,
        max_iter=8,
    )

    def cb(frac: float, msg: str):
        return lambda _out: lg.step(base_pct + int(span * frac), msg)

    t1 = Task(
        description=(
            "summarize_voc_by_industry 도구를 한 번 호출하여 산업군별 VOC 데이터를 확인하고, "
            "산업군마다 어떤 VOC 가 주로 접수되는지 건수/비율/분야/키워드를 근거로 정리하세요."
        ),
        expected_output="산업군별 주요 VOC 분석 요약 (산업군마다 2~3문장)",
        agent=agent,
        callback=cb(0.33, "[Issue Agent] 1/3 산업군별 주요 VOC 분석 완료"),
    )
    t2 = Task(
        description=(
            "앞선 분석을 바탕으로 산업군별 주요 이슈를 2~3개씩 도출하고, 산업군 공통 이슈도 정리하세요. "
            "각 이슈에는 근거(비율, 키워드, 사례)를 함께 적습니다."
        ),
        expected_output="산업군별 주요 이슈 목록과 산업군 공통 이슈",
        agent=agent,
        context=[t1],
        callback=cb(0.66, "[Issue Agent] 2/3 주요 이슈 도출 완료"),
    )
    t3 = Task(
        description=(
            "도출된 이슈마다 실행 가능한 개선과제를 산업군별 2~3개씩 제시하고 우선순위(상/중/하)를 매기세요. "
            "분석 대상 산업군을 하나도 빠뜨리지 말고 최종 결과를 지정된 JSON 구조로 출력하세요."
        ),
        expected_output="산업군별 주요 VOC 요약, 이슈, 개선과제와 전체 총평, 공통 이슈를 담은 구조화 결과",
        agent=agent,
        context=[t1, t2],
        output_pydantic=IssueAnalysis,
        callback=cb(1.0, "[Issue Agent] 3/3 개선과제 도출 완료"),
    )
    lg.step(base_pct, "[Issue Agent] CrewAI 실행: 산업군별 VOC 분석 -> 이슈 도출 -> 개선과제 도출")
    run_crew(agent, [t1, t2, t3])
    if "summarize_voc_by_industry" not in ctx.called:
        lg.log("[경고] Issue Agent 가 데이터 요약 도구를 호출하지 않았습니다. 결과를 확인해 주세요.")
    return parse_model(t3, IssueAnalysis)


# ── Report Agent 실행 ────────────────────────────────────────────────────────
def issues_markdown(issues: IssueAnalysis) -> str:
    lines = [f"**총평**: {issues.overall_summary}", ""]
    for it in issues.industries:
        lines += [f"#### {it.industry}", it.voc_summary, "", "**주요 이슈**"]
        lines += [f"- {x}" for x in it.issues]
        lines += ["", "**개선과제**"]
        lines += [f"- [{x.priority}] {x.action}" for x in it.improvements]
        lines.append("")
    if issues.common_issues:
        lines += ["#### 산업군 공통 이슈"] + [f"- {x}" for x in issues.common_issues]
    return "\n".join(lines)


def report_markdown(content: ReportContent, issues: IssueAnalysis) -> str:
    plan = ["| 우선순위 | 산업군 | 이슈 | 대응방안 | 일정 | 기대효과 |", "|---|---|---|---|---|---|"]
    for a in content.action_plan:
        plan.append(f"| {a.priority} | {a.industry} | {a.issue} | {a.action} | {a.timeline} | {a.effect} |")
    return "\n".join(
        [
            f"### {content.title}",
            "**핵심 요약**",
            content.executive_summary,
            "",
            "### 주요 이슈",
            issues_markdown(issues),
            "",
            "### 대응방안",
            "\n".join(plan),
            "",
            "**결론**: " + content.conclusion,
        ]
    )


def run_report_agent(ctx: VocContext, lg: JobLog, base_pct: int, span: int) -> str:
    assert ctx.issues is not None
    wc = ctx.wc.get("전체")
    brief = no_braces(
        "\n".join(
            [
                f"[개요] 파일 {ctx.filename}, 총 {len(ctx.df):,}건, 기간 {period_text(ctx.df)}",
                stats_brief(ctx, "산업군"),
                stats_brief(ctx, "분야"),
                "[산업군별 상위 분야] " + cross_brief(ctx),
                wc.text if wc else "",
                "[Issue Agent 분석 결과]\n" + issues_markdown(ctx.issues),
            ]
        )
    )

    def create_word_report_impl() -> str:
        if ctx.report_content is None:
            return "오류: 보고서 내용이 아직 작성되지 않았습니다."
        lg.log("[Report Agent] 도구 실행 - Word 보고서 파일 생성")
        ctx.called.add("create_word_report")
        ctx.report_path = build_word_report(ctx, ctx.report_content, lg)
        return f"보고서 파일 생성 완료: {ctx.report_path}"

    @tool("create_word_report")
    def create_word_report() -> str:
        """작성된 보고서 내용과 VOC 통계, 그래프, 워드클라우드, 이슈/대응방안을 담은 Word(.docx) 파일을 생성하고 파일 경로를 반환합니다."""
        return create_word_report_impl()

    agent = Agent(
        role="VOC 보고서 작성 전문가",
        goal="VOC 분석 결과를 경영진이 의사결정에 쓸 수 있는 Word 보고서로 만든다",
        backstory="임원 보고서를 다수 작성한 컨설턴트. 결론 먼저 쓰고, 숫자 근거를 제시하며, 간결한 한국어 문장을 사용한다.",
        tools=[create_word_report],
        llm=get_llm(),
        verbose=False,
        allow_delegation=False,
        max_iter=6,
    )

    def store_content(out) -> None:
        try:
            ctx.report_content = out.pydantic or ReportContent.model_validate_json(re.search(r"\{.*\}", out.raw, re.S).group(0))
        except Exception:  # noqa: BLE001
            ctx.report_content = None
        lg.step(base_pct + int(span * 0.6), "[Report Agent] 1/2 보고서 내용 작성 완료")

    t1 = Task(
        description=(
            "아래 분석 자료만을 근거로 VOC 분석 보고서 내용을 작성하세요. 수치는 자료에 있는 값만 사용하고 비율은 소수점 둘째 자리까지 표기합니다. "
            "대응방안(action_plan)은 Issue Agent 의 개선과제를 구체화해 6~10개로 작성하고 우선순위 순으로 정렬하세요.\n\n"
            + brief
        ),
        expected_output="제목, 핵심 요약, 산업군/분야/키워드 해설, 대응방안 표, 결론을 담은 구조화 결과",
        agent=agent,
        output_pydantic=ReportContent,
        callback=store_content,
    )
    t2 = Task(
        description="create_word_report 도구를 정확히 한 번 호출하여 Word 보고서 파일을 생성하고, 생성된 파일 경로를 알려주세요.",
        expected_output="생성된 Word 파일 경로",
        agent=agent,
        context=[t1],
    )
    lg.step(base_pct, "[Report Agent] CrewAI 실행: 보고서 작성 -> Word 파일 생성")
    run_crew(agent, [t1, t2])
    if ctx.report_content is None:
        raise UserError("Report Agent 가 보고서 내용을 만들지 못했습니다. 다시 시도해 주세요.")
    if not ctx.report_path or not Path(ctx.report_path).exists():
        lg.log("[Report Agent] 도구 호출이 누락되어 보고서 파일을 직접 생성합니다.")
        ctx.report_path = build_word_report(ctx, ctx.report_content, lg)
    return ctx.report_path


# ════════════════════════════════════════════════════════════════════════════
# 6. Word 보고서 생성 (한글 폰트 보장)
# ════════════════════════════════════════════════════════════════════════════
def _set_rfonts(rpr, name: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if qn(attr) in rfonts.attrib:
            del rfonts.attrib[qn(attr)]


def apply_korean_fonts(doc: Document, name: str = DOCX_FONT) -> None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None and rpr_default.find(qn("w:rPr")) is not None:
            _set_rfonts(rpr_default.find(qn("w:rPr")), name)
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            style.font.name = name
            _set_rfonts(style.element.get_or_add_rPr(), name)


def shade_cell(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc: Document, header: list[str], rows: list[list[str]], widths_cm: list[float] | None = None, size: int = 9) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(cell, "3B4A8C")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(size)
    if widths_cm:
        for r in table.rows:
            for i, w in enumerate(widths_cm):
                r.cells[i].width = Cm(w)
    doc.add_paragraph()


def add_text_block(doc: Document, text: str) -> None:
    for line in [ln.strip() for ln in str(text).splitlines() if ln.strip()]:
        if line[0] in "-•·*":
            doc.add_paragraph(line.lstrip("-•·* ").strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)


def save_bar_png(res: StatResult, dim: str, path: Path) -> None:
    """Plotly(kaleido) 로 저장하고, 실패하면 matplotlib 으로 대체한다."""
    try:
        res.fig.write_image(str(path), width=1000, height=480, scale=2)
        return
    except Exception:  # noqa: BLE001
        pass
    from matplotlib import font_manager

    if FONT_PATH:
        font_manager.fontManager.addfont(FONT_PATH)
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=FONT_PATH).get_name()
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(10, 4.8))
    bars = ax.bar(res.table["항목"], res.table["비율(%)"], color="#4f46e5")
    ax.bar_label(bars, labels=[f"{v:.2f}%" for v in res.table["비율(%)"]])
    ax.set_ylabel("비율(%)")
    ax.set_title(f"{dim}별 VOC 비율")
    plt.xticks(rotation=30, ha="right")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def cross_table_rows(ctx: VocContext, max_fields: int = 8) -> tuple[list[str], list[list[str]]]:
    df = ctx.df
    top_fields = list(df["분야"].value_counts().head(max_fields).index)
    d = df.assign(분야=df["분야"].where(df["분야"].isin(top_fields), "기타"))
    cols = top_fields + (["기타"] if (d["분야"] == "기타").any() else [])
    ct = pd.crosstab(d["산업군"], d["분야"]).reindex(columns=cols, fill_value=0)
    ct = ct.loc[ct.sum(axis=1).sort_values(ascending=False).index]
    header = ["산업군"] + cols + ["합계"]
    rows = [[ind] + [f"{int(v):,}" for v in r] + [f"{int(r.sum()):,}"] for ind, r in ct.iterrows()]
    totals = ct.sum(axis=0)
    rows.append(["합계"] + [f"{int(v):,}" for v in totals] + [f"{int(totals.sum()):,}"])
    return header, rows


def build_word_report(ctx: VocContext, content: ReportContent, lg: JobLog | None = None) -> str:
    log = lg.log if lg else (lambda _m: None)
    issues = ctx.issues
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    apply_korean_fonts(doc)
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x5A)
    doc.core_properties.title = content.title

    log("[Report Agent] 표지/개요 작성")
    doc.add_heading(content.title, level=0)
    p = doc.add_paragraph(f"작성일: {datetime.now():%Y-%m-%d}   |   분석 파일: {ctx.filename}")
    p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    df = ctx.df
    doc.add_heading("1. 분석 개요", level=1)
    add_table(
        doc,
        ["구분", "내용"],
        [
            ["총 VOC 건수", f"{len(df):,}건"],
            ["분석 기간", period_text(df)],
            ["산업군 / 제품 / 분야", f"{df['산업군'].nunique()}개 / {df['제품명'].nunique()}개 / {df['분야'].nunique()}개"],
            ["고객사 수", f"{df['고객명'].nunique()}곳"],
        ],
        [5, 11],
        size=10,
    )
    doc.add_heading("핵심 요약", level=2)
    add_text_block(doc, content.executive_summary)

    for no, dim, comment in ((2, "산업군", content.industry_comment), (3, "분야", content.field_comment)):
        log(f"[Report Agent] {dim}별 통계 섹션 작성")
        if dim not in ctx.stats:
            compute_stats(ctx, dim)
        res = ctx.stats[dim]
        doc.add_heading(f"{no}. {dim}별 VOC 통계", level=1)
        img = ctx.workdir / f"chart_{dim}.png"
        save_bar_png(res, dim, img)
        doc.add_picture(str(img), width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_table(
            doc,
            [dim, "건수", "비율(%)"],
            [[r.항목, f"{r.건수:,}", f"{r._4:.2f}%"] for r in res.table.itertuples()],
            [7, 4, 4],
        )
        add_text_block(doc, comment)
        if dim == "분야":
            doc.add_heading("산업군 x 분야 교차 건수", level=2)
            header, rows = cross_table_rows(ctx)
            add_table(doc, header, rows, size=8)

    log("[Report Agent] 워드클라우드 섹션 작성")
    doc.add_heading("4. 불만 키워드 분석 (워드클라우드)", level=1)
    wc = ctx.wc.get("전체")
    if wc:
        doc.add_picture(wc.path, width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_table(doc, ["순위", "키워드", "언급 횟수"], [[i + 1, w, f"{c:,}"] for i, (w, c) in enumerate(wc.keywords[:15])], [3, 7, 5])
    add_text_block(doc, content.wordcloud_comment)

    log("[Report Agent] 주요 이슈/대응방안 섹션 작성")
    doc.add_heading("5. 주요 이슈", level=1)
    if issues:
        add_text_block(doc, issues.overall_summary)
        for it in issues.industries:
            doc.add_heading(it.industry, level=2)
            add_text_block(doc, it.voc_summary)
            for x in it.issues:
                doc.add_paragraph(x, style="List Bullet")
        if issues.common_issues:
            doc.add_heading("산업군 공통 이슈", level=2)
            for x in issues.common_issues:
                doc.add_paragraph(x, style="List Bullet")

    doc.add_heading("6. 대응방안", level=1)
    add_table(
        doc,
        ["우선순위", "산업군", "이슈", "대응방안", "일정", "기대효과"],
        [[a.priority, a.industry, a.issue, a.action, a.timeline, a.effect] for a in content.action_plan],
        [1.5, 2, 3.4, 4.6, 1.8, 3],
        size=8.5,
    )
    doc.add_heading("7. 결론", level=1)
    add_text_block(doc, content.conclusion)

    path = ctx.workdir / f"VOC_분석보고서_{datetime.now():%Y%m%d_%H%M%S}.docx"
    doc.save(str(path))
    log(f"[Report Agent] 보고서 저장: {path.name}")
    return str(path)


# ════════════════════════════════════════════════════════════════════════════
# 7. Gradio UI
# ════════════════════════════════════════════════════════════════════════════
def need_ctx(ctx: VocContext | None) -> VocContext:
    if ctx is None:
        raise UserError("먼저 '파일 업로드' 메뉴에서 CSV 파일을 업로드해 주세요.")
    return ctx


def upload_outputs(ctx: VocContext):
    industries = ["전체"] + sorted(ctx.df["산업군"].unique())
    return (
        ctx,
        ctx.df,
        overview_md(ctx),
        gr.update(choices=industries, value="전체"),
    )


def on_upload(path, ctx):
    def work(lg: JobLog):
        if not path:
            raise UserError("CSV 파일을 선택해 주세요.")
        return load_voc_csv(path, lg)

    yield from stream_job(work, upload_outputs, 4)


def on_sample(ctx):
    def work(lg: JobLog):
        if not SAMPLE_CSV.exists():
            raise UserError("sample_voc.csv 파일이 없습니다.")
        return load_voc_csv(str(SAMPLE_CSV), lg)

    yield from stream_job(work, upload_outputs, 4)


def on_stats(dims, ctx):
    def work(lg: JobLog):
        c = need_ctx(ctx)
        if not dims:
            raise UserError("통계 항목(산업군/제품명/분야)을 한 개 이상 선택해 주세요.")
        selected = [d for d in STAT_DIMS if d in dims]
        lg.step(5, f"[VOC Agent] 통계분석 시작: {', '.join(selected)}")
        return selected, run_voc_stats(c, selected, lg)

    def final(res):
        selected, insight = res
        c = ctx
        plots = [
            gr.update(value=c.stats[d].fig, visible=True) if d in selected else gr.update(value=None, visible=False)
            for d in STAT_DIMS
        ]
        table = format_stats_table([c.stats[d] for d in selected])
        return [insight, *plots, gr.update(value=table, visible=True)]

    yield from stream_job(work, final, 5)


def on_wordcloud(industry, max_words, stop_text, ctx):
    def work(lg: JobLog):
        c = need_ctx(ctx)
        extra = [s for s in re.split(r"[,\s]+", stop_text or "") if s]
        lg.step(5, f"[VOC Agent] 워드클라우드 시작 (산업군: {industry}, 최대 {int(max_words)}단어)")
        return run_voc_wordcloud(c, {"industry": industry, "max_words": int(max_words), "extra_stop": extra}, lg)

    def final(res):
        wc, insight = res
        return [wc.path, keyword_bar(wc), insight]

    yield from stream_job(work, final, 3)


def on_report(ctx):
    def work(lg: JobLog):
        c = need_ctx(ctx)
        require_api_key()
        c.called.clear()
        c.issues = c.report_content = c.report_path = None
        lg.step(3, "[Report] 보고서 생성 파이프라인 시작 (VOC -> Issue -> Report)")

        lg.step(5, "[VOC Agent] 산업군/분야 통계 + 워드클라우드 생성")
        run_voc_stats(c, ["산업군", "분야"], lg)
        run_voc_wordcloud(c, {"industry": "전체", "max_words": 80, "extra_stop": []}, lg)

        lg.step(35, "[Issue Agent] 시작")
        c.issues = run_issue_agent(c, lg, base_pct=35, span=35)

        lg.step(72, "[Report Agent] 시작")
        path = run_report_agent(c, lg, base_pct=72, span=25)
        lg.step(98, "[Report Agent] 다운로드 준비 완료")
        return path

    def final(path):
        return [path, report_markdown(ctx.report_content, ctx.issues)]

    yield from stream_job(work, final, 2)


CSS = """
.gradio-container {font-family: 'Malgun Gothic','맑은 고딕','Apple SD Gothic Neo',sans-serif !important;}
#log-box textarea {font-family: Consolas,'Malgun Gothic',monospace; font-size: 12.5px;}
"""


def build_ui() -> gr.Blocks:
    with gr.Blocks(title="고객사 VOC 분석 에이전트") as demo:
        ctx_state = gr.State(None)
        gr.Markdown(
            "# 고객사 VOC 분석 에이전트\n"
            "CrewAI 기반 **VOC Agent · Issue Agent · Report Agent** 가 VOC 통계/키워드/이슈/보고서를 자동 생성합니다."
        )

        with gr.Tabs():
            with gr.Tab("파일 업로드"):
                gr.Markdown(f"CSV 컬럼: `{', '.join(REQUIRED_COLUMNS)}`")
                with gr.Row():
                    file_in = gr.File(label="VOC CSV 업로드", file_types=[".csv"], file_count="single", scale=3)
                    sample_btn = gr.Button("샘플 데이터 불러오기", scale=1, visible=SAMPLE_CSV.exists())
                info_md = gr.Markdown()
                data_df = gr.Dataframe(label="업로드 데이터", interactive=False, wrap=True, max_height=420)

            with gr.Tab("통계분석"):
                dims_in = gr.CheckboxGroup(STAT_DIMS, value=["산업군"], label="분석 항목 선택", info="선택한 항목별 비율(%) 막대그래프를 표시합니다.")
                stats_btn = gr.Button("통계분석 실행", variant="primary")
                stats_md = gr.Markdown()
                plots = [gr.Plot(visible=False, label=f"{d}별 비율") for d in STAT_DIMS]
                stats_table = gr.Dataframe(label="통계 표", interactive=False, visible=False, max_height=380)

            with gr.Tab("워드클라우드"):
                with gr.Row():
                    wc_industry = gr.Dropdown(["전체"], value="전체", label="산업군", scale=1)
                    wc_max = gr.Slider(20, 200, value=80, step=10, label="최대 단어 수", scale=1)
                    wc_stop = gr.Textbox(label="제외할 단어 (쉼표/공백 구분)", placeholder="예: 제품, 서비스", scale=2)
                wc_btn = gr.Button("워드클라우드 생성", variant="primary")
                wc_img = gr.Image(label="워드클라우드", interactive=False, type="filepath")
                wc_md = gr.Markdown()
                kw_plot = gr.Plot(label="상위 키워드")

            with gr.Tab("보고서생성"):
                gr.Markdown("VOC Agent(통계/워드클라우드) → Issue Agent(이슈/개선과제) → Report Agent(Word 보고서) 순서로 실행됩니다. 1~3분 정도 걸릴 수 있습니다.")
                rep_btn = gr.Button("보고서 생성", variant="primary")
                rep_file = gr.File(label="보고서 다운로드 (.docx)", interactive=False)
                rep_md = gr.Markdown()

        with gr.Group():
            prog = gr.HTML(progress_html(0, "대기 중"))
            log_box = gr.Textbox(label="실행 로그", lines=12, max_lines=12, interactive=False, autoscroll=True, elem_id="log-box")

        common = [prog, log_box]
        up_out = common + [ctx_state, data_df, info_md, wc_industry]
        file_in.upload(on_upload, [file_in, ctx_state], up_out)
        sample_btn.click(on_sample, [ctx_state], up_out)
        stats_btn.click(on_stats, [dims_in, ctx_state], common + [stats_md, *plots, stats_table])
        wc_btn.click(on_wordcloud, [wc_industry, wc_max, wc_stop, ctx_state], common + [wc_img, kw_plot, wc_md])
        rep_btn.click(on_report, [ctx_state], common + [rep_file, rep_md])
    return demo


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(exist_ok=True)
    print(f"한글 폰트: {FONT_PATH} | LLM: {LLM_MODEL} | API KEY: {'설정됨' if os.getenv('OPENAI_API_KEY') else '없음'}")
    build_ui().launch(
        inbrowser=True,
        theme=gr.themes.Soft(primary_hue="indigo"),
        css=CSS,
        allowed_paths=[str(OUTPUT_DIR)],
    )
