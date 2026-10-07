"""뉴스 수집·분석 에이전트 (LangGraph + Gradio + Plotly + SerpAPI)

실행:  python news_agent_app.py

그래프 노드
  news    : SerpAPI 로 구글 뉴스 / 네이버 뉴스 수집
  analyze : 이슈 그룹핑, 요약·시사점, 워드클라우드, 빈도 그래프
  report  : Word(.docx) 보고서 생성 후 다운로드

사전 조건
  - .env 의 SERPAPI_API_KEY
  - 환경변수 OPENAI_API_KEY
"""

from __future__ import annotations

import html
import json
import os
import re
import sys
import threading
import time
import traceback
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, TypedDict
from urllib.parse import urlparse

from dotenv import load_dotenv

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import warnings

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

warnings.filterwarnings("ignore", message="function callbacks cannot be serialized")

import gradio as gr
import pandas as pd
import plotly.graph_objects as go
import requests
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from langchain_openai import ChatOpenAI
from langgraph.graph import END, START, StateGraph
from pydantic import BaseModel, Field, field_validator
from wordcloud import WordCloud

# ════════════════════════════════════════════════════════════════════════════
# 1. 설정 / 한글 폰트
# ════════════════════════════════════════════════════════════════════════════
OUTPUT_DIR = BASE_DIR / "outputs"
LLM_MODEL = os.getenv("NEWS_LLM_MODEL", "gpt-4o-mini")
SERPAPI_URL = "https://serpapi.com/search.json"
MAX_PAGES = 10
REQUEST_TIMEOUT = 45

PLOTLY_FONT = "Malgun Gothic, 맑은 고딕, NanumGothic, Apple SD Gothic Neo, Noto Sans CJK KR, sans-serif"
DOCX_FONT = "맑은 고딕"

INK = "#12323C"
TEAL = "#0E7C7B"
GOLD = "#C9A227"
COPPER = "#B85C38"
PAPER = "#F2F5F3"


def find_korean_font() -> str | None:
    """워드클라우드/matplotlib 용 한글 폰트 파일 경로."""
    candidates = [
        r"C:\Windows\Fonts\malgun.ttf",
        r"C:\Windows\Fonts\malgunbd.ttf",
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/Library/Fonts/NanumGothic.ttf",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    from matplotlib import font_manager

    keys = ("malgun", "nanum", "gothic", "noto sans cjk kr", "apple sd")
    for f in font_manager.fontManager.ttflist:
        if any(k in f.name.lower() for k in keys):
            return f.fname
    return None


FONT_PATH = find_korean_font()


def setup_matplotlib_korean() -> None:
    from matplotlib import font_manager

    if FONT_PATH:
        try:
            font_manager.fontManager.addfont(FONT_PATH)
        except Exception:
            pass
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=FONT_PATH).get_name()
    plt.rcParams["axes.unicode_minus"] = False


setup_matplotlib_korean()


class UserError(Exception):
    """사용자에게 그대로 보여줄 오류."""


def serpapi_key() -> str:
    key = (os.getenv("SERPAPI_API_KEY") or os.getenv("SERPAPI_KEY") or "").strip()
    if not key:
        raise UserError("`.env` 파일에 SERPAPI_API_KEY 를 입력해 주세요.")
    return key


def require_openai() -> None:
    if not os.getenv("OPENAI_API_KEY"):
        raise UserError("환경변수 OPENAI_API_KEY 가 설정되어 있지 않습니다.")


# ════════════════════════════════════════════════════════════════════════════
# 2. 진행 로그
# ════════════════════════════════════════════════════════════════════════════
class JobLog:
    def __init__(self) -> None:
        self.lines: list[str] = []
        self.pct = 0
        self.status = "대기 중"
        self.failed = False
        self.version = 0
        self._lock = threading.Lock()

    def log(self, msg: str) -> None:
        with self._lock:
            self.lines.append(f"[{datetime.now():%H:%M:%S}] {msg}")
            self.version += 1

    def step(self, pct: int, msg: str) -> None:
        self.pct, self.status = pct, msg
        self.log(msg)

    def fail(self, msg: str) -> None:
        self.failed = True
        self.status = "오류 발생"
        self.log(f"[오류] {msg}")

    def text(self) -> str:
        with self._lock:
            return "\n".join(self.lines)


def progress_html(pct: int, status: str, failed: bool = False) -> str:
    color = "#c45c38" if failed else TEAL
    return (
        '<div style="margin:2px 0 8px 0;">'
        f'<div style="display:flex;justify-content:space-between;font-size:13px;margin-bottom:4px;color:{INK};">'
        f"<span><b>진행</b> · {status}</span><span>{pct}%</span></div>"
        f'<div style="background:#d7e0e4;border-radius:2px;height:10px;overflow:hidden;">'
        f'<div style="width:{pct}%;height:100%;background:{color};transition:width .3s;"></div>'
        "</div></div>"
    )


def stream_job(work, final_fn, n_extra: int):
    lg = JobLog()
    holder: dict[str, Any] = {}

    def runner():
        try:
            holder["res"] = work(lg)
        except UserError as e:
            lg.fail(str(e))
        except Exception as e:  # noqa: BLE001
            lg.fail(f"{type(e).__name__}: {e}")
            lg.log(traceback.format_exc(limit=4))

    noop = [gr.update()] * n_extra
    thread = threading.Thread(target=runner, daemon=True)
    thread.start()

    def head():
        return [progress_html(lg.pct, lg.status, lg.failed), lg.text()]

    yield head() + noop
    last = -1
    while thread.is_alive():
        if lg.version != last:
            last = lg.version
            yield head() + noop
        time.sleep(0.3)
    thread.join()

    if lg.failed or "res" not in holder:
        yield head() + noop
        return
    lg.pct, lg.status = 100, "완료"
    lg.log("완료")
    yield head() + list(final_fn(holder["res"]))


# ════════════════════════════════════════════════════════════════════════════
# 3. SerpAPI 뉴스 수집
# ════════════════════════════════════════════════════════════════════════════
def clean_text(value: Any) -> str:
    text = html.unescape(str(value or ""))
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_date(value: Any) -> str:
    raw = clean_text(value)
    if not raw:
        return ""
    iso = re.match(r"(\d{4}-\d{2}-\d{2})", raw)
    if iso:
        return iso.group(1)
    mdy = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", raw)
    if mdy:
        month, day, year = mdy.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"
    ymd = re.match(r"(\d{4})\.(\d{1,2})\.(\d{1,2})", raw)
    if ymd:
        year, month, day = ymd.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"
    return raw


def source_name(raw: Any) -> str:
    if isinstance(raw, dict):
        return clean_text(raw.get("name") or raw.get("title") or "")
    return clean_text(raw)


def news_url(raw: dict) -> str:
    return clean_text(raw.get("link") or raw.get("news_link") or raw.get("url") or "")


def serpapi_search(params: dict[str, Any]) -> dict[str, Any]:
    payload = {**params, "api_key": serpapi_key(), "output": "json"}
    try:
        resp = requests.get(SERPAPI_URL, params=payload, timeout=REQUEST_TIMEOUT)
        resp.raise_for_status()
        data = resp.json()
    except requests.RequestException as e:
        raise UserError(f"SerpAPI 요청에 실패했습니다: {e}") from e
    if error := data.get("error"):
        raise UserError(f"SerpAPI 오류: {error}")
    return data


def parse_news_item(raw: dict) -> dict[str, str] | None:
    title = clean_text(raw.get("title"))
    url = news_url(raw)
    if not title or not url:
        return None
    info = raw.get("news_info") if isinstance(raw.get("news_info"), dict) else {}
    date = (
        raw.get("iso_date")
        or raw.get("date")
        or raw.get("news_date")
        or info.get("news_date")
        or ""
    )
    snippet = clean_text(raw.get("snippet") or raw.get("snippet_highlighted_words") or "")
    if isinstance(raw.get("snippet_highlighted_words"), list) and not snippet:
        snippet = clean_text(" ".join(str(x) for x in raw["snippet_highlighted_words"]))
    source = (
        source_name(raw.get("source"))
        or clean_text(raw.get("press_name"))
        or clean_text(info.get("press_name"))
    )
    return {
        "날짜": normalize_date(date),
        "제목": title,
        "주요내용": snippet or title,
        "url": url,
        "출처": source,
    }


def iter_google_entries(results: list[dict]) -> list[dict]:
    entries: list[dict] = []
    for item in results:
        if item.get("title") and (item.get("link") or item.get("news_link")):
            entries.append(item)
        highlight = item.get("highlight")
        if isinstance(highlight, dict):
            entries.append(highlight)
        for story in item.get("stories") or []:
            if isinstance(story, dict):
                entries.append(story)
    return entries


def append_unique(bucket: list[dict], seen: set[str], raw: dict) -> None:
    item = parse_news_item(raw)
    if not item:
        return
    key = item["url"].split("?")[0].rstrip("/").lower()
    host_title = (urlparse(item["url"]).netloc, item["제목"])
    if key in seen or host_title in seen:
        return
    seen.add(key)
    seen.add(host_title)  # type: ignore[arg-type]
    bucket.append(item)


def fetch_google_news(keyword: str, count: int, lg: JobLog) -> list[dict]:
    items: list[dict] = []
    seen: set[str] = set()

    lg.log("[news] 구글 뉴스(engine=google_news) 요청")
    data = serpapi_search({"engine": "google_news", "q": keyword, "gl": "kr", "hl": "ko"})
    for raw in iter_google_entries(data.get("news_results") or []):
        append_unique(items, seen, raw)
        if len(items) >= count:
            return number_items(items[:count])
    lg.log(f"[news] google_news {len(items)}건. 부족분이 있으면 Google 뉴스 탭으로 보충합니다.")

    start = 0
    page = 0
    while len(items) < count and page < MAX_PAGES:
        page += 1
        lg.log(f"[news] Google 뉴스 탭 보충 {page}페이지")
        data = serpapi_search(
            {
                "engine": "google",
                "q": keyword,
                "tbm": "nws",
                "hl": "ko",
                "gl": "kr",
                "num": 10,
                "start": start,
                "tbs": "sbd:1",
            }
        )
        batch = data.get("news_results") or []
        if not batch:
            break
        for raw in batch:
            append_unique(items, seen, raw)
            if len(items) >= count:
                break
        start += 10
    return number_items(items[:count])


def fetch_naver_news(keyword: str, count: int, lg: JobLog) -> list[dict]:
    items: list[dict] = []
    seen: set[str] = set()
    page = 1
    while len(items) < count and page <= MAX_PAGES:
        lg.log(f"[news] 네이버 뉴스 {page}페이지 요청")
        data = serpapi_search(
            {
                "engine": "naver",
                "query": keyword,
                "where": "news",
                "sort_by": "1",
                "page": page,
            }
        )
        batch = data.get("news_results") or []
        if not batch:
            break
        for raw in batch:
            append_unique(items, seen, raw)
            if len(items) >= count:
                break
        page += 1
    return number_items(items[:count])


def number_items(items: list[dict]) -> list[dict]:
    numbered = []
    for i, item in enumerate(items, start=1):
        numbered.append({"순번": i, **item})
    return numbered


def collect_news(keyword: str, count: int, source: str, lg: JobLog) -> list[dict]:
    if source == "네이버 뉴스":
        items = fetch_naver_news(keyword, count, lg)
    else:
        items = fetch_google_news(keyword, count, lg)
    if not items:
        raise UserError(f"'{keyword}' 관련 뉴스를 찾지 못했습니다. 키워드를 바꿔 보세요.")
    lg.log(f"[news] 수집 완료 {len(items)}건 (요청 {count}건)")
    return items


# ════════════════════════════════════════════════════════════════════════════
# 4. 키워드 / 워드클라우드 / 그래프
# ════════════════════════════════════════════════════════════════════════════
KO_STOPWORDS = set(
    "너무 정말 매우 계속 자꾸 항상 여러 대한 통해 관련 경우 때문 이후 현재 최근 또한 그리고 하지만 그러나 "
    "있는 없는 있어 없어 있음 없음 같은 같이 많이 아주 조금 다시 또는 이런 저희 우리 지속 발생 요청 문의 "
    "진행 확인 필요 사용 제공 문제 부분 정도 대비 이상 이하 모든 해당 가장 매번 자주 하는 하고 해서 "
    "있습니다 없습니다 합니다 입니다 사항 내용 상황 대해 위해 따라 로서 에서 으로 에게 그냥 대부분 일부 "
    "전혀 다른 않아 않고 않는 않은 않았 못해 못한 기자 뉴스 보도 사진 영상 오늘 어제 지난 올해 내년 "
    "한국 대한민국 정부 서울 연합뉴스 따르면 밝혔다 전했다 이다 있다 했다 하며 이번 지난달 이달 오후 "
    "오전 지난해 올해부터 가운데 전망 분석 업계 시장 기업 국내 해외 세계 글로벌 발표 계획 추진".split()
)
KO_ENDINGS = sorted(
    "했습니다 합니다 됩니다 입니다 습니다 에서는 에서도 으로는 으로도 에서 으로 에게 까지 부터 보다 "
    "처럼 마다 이나 이며 이고 이다 하고 해서 하여 하는 하지 하게 되어 되는 되고 돼서 됨 함 에는 에도 "
    "에만 이라고 이라 라고 한다 된다 인데 지만 으며 거나 였다 했다 됐다 적인 적으로 한 된".split(),
    key=len,
    reverse=True,
)
KO_JOSA_STRONG = set("을를은는이가에")
KO_JOSA_WEAK = set("의도로과와만")
KO_PROTECTED = {"전문가", "사업가", "투자가", "신뢰도", "만족도", "인지도"}
TOKEN_RE = re.compile(r"[가-힣]{2,}|[A-Za-z][A-Za-z0-9\-]{1,}")


def extract_keywords(texts, extra_stop=()) -> Counter:
    texts = [str(t) for t in texts if str(t).strip()]
    raw = Counter(tok for t in texts for tok in TOKEN_RE.findall(t))
    stop = KO_STOPWORDS | {s.strip() for s in extra_stop if s.strip()}

    def normalize(tok: str) -> str | None:
        if not re.fullmatch(r"[가-힣]+", tok):
            return tok.upper() if len(tok) <= 6 else tok
        if tok.endswith(("요", "다", "까", "죠")):
            return None
        for end in KO_ENDINGS:
            if tok.endswith(end) and len(tok) - len(end) >= 2:
                return tok[: -len(end)]
        if len(tok) >= 3 and tok not in KO_PROTECTED:
            if tok[-1] in KO_JOSA_STRONG or (tok[-1] in KO_JOSA_WEAK and tok[:-1] in raw):
                return tok[:-1]
        return tok

    result: Counter = Counter()
    for tok, n in raw.items():
        word = normalize(tok)
        if word and len(word) >= 2 and word not in stop and tok not in stop:
            result[word] += n
    return Counter({k: v for k, v in result.items() if v >= 1})


def news_corpus(items: list[dict]) -> list[str]:
    return [f"{it.get('제목', '')} {it.get('주요내용', '')}" for it in items]


def build_wordcloud(freq: list[tuple[str, int]], path: Path, max_words: int = 80) -> None:
    if not FONT_PATH:
        raise UserError("한글 폰트를 찾을 수 없습니다. 맑은 고딕 또는 나눔고딕을 설치해 주세요.")
    if not freq:
        raise UserError("워드클라우드를 그릴 키워드가 없습니다.")
    palette = [INK, TEAL, GOLD, COPPER, "#2A6F7B", "#4A7C59"]

    def color_func(*_args, **_kwargs):
        import random

        return random.choice(palette)

    wc = WordCloud(
        font_path=FONT_PATH,
        width=1200,
        height=620,
        background_color="white",
        color_func=color_func,
        prefer_horizontal=0.92,
        max_font_size=140,
        relative_scaling=0.42,
        margin=4,
        max_words=int(max_words),
    ).generate_from_frequencies(dict(freq[:max_words]))
    wc.to_file(str(path))


def style_fig(fig: go.Figure, title: str, height: int = 430) -> go.Figure:
    fig.update_layout(
        title=dict(text=title, font=dict(size=17, color=INK, family=PLOTLY_FONT)),
        template="plotly_white",
        paper_bgcolor="white",
        plot_bgcolor="white",
        font=dict(family=PLOTLY_FONT, size=13, color=INK),
        height=height,
        margin=dict(l=80, r=30, t=70, b=50),
        showlegend=False,
    )
    fig.update_xaxes(title_font=dict(family=PLOTLY_FONT), tickfont=dict(family=PLOTLY_FONT))
    fig.update_yaxes(title_font=dict(family=PLOTLY_FONT), tickfont=dict(family=PLOTLY_FONT))
    return fig


def keyword_bar(freq: list[tuple[str, int]], n: int = 20) -> go.Figure:
    top = freq[:n][::-1]
    fig = go.Figure(
        go.Bar(
            x=[c for _, c in top],
            y=[w for w, _ in top],
            orientation="h",
            text=[f"{c:,}" for _, c in top],
            textposition="outside",
            marker_color=TEAL,
            hovertemplate="%{y}: %{x:,}회<extra></extra>",
        )
    )
    style_fig(fig, "주요 단어 발생 빈도", height=max(380, 22 * len(top) + 120))
    fig.update_xaxes(title_text="발생 횟수")
    fig.update_yaxes(title_text="")
    return fig


def save_keyword_png(freq: list[tuple[str, int]], path: Path, n: int = 20) -> None:
    fig = keyword_bar(freq, n)
    try:
        fig.write_image(str(path), width=1000, height=max(480, 28 * min(n, len(freq)) + 140), scale=2)
        return
    except Exception:
        pass
    top = freq[:n][::-1]
    setup_matplotlib_korean()
    fig_m, ax = plt.subplots(figsize=(10, max(4.5, 0.32 * len(top) + 1.6)))
    ax.barh([w for w, _ in top], [c for _, c in top], color=TEAL)
    ax.set_xlabel("발생 횟수")
    ax.set_title("주요 단어 발생 빈도")
    fig_m.tight_layout()
    fig_m.savefig(path, dpi=150)
    plt.close(fig_m)


# ════════════════════════════════════════════════════════════════════════════
# 5. LangGraph 상태 / LLM 분석
# ════════════════════════════════════════════════════════════════════════════
class NewsGroup(BaseModel):
    title: str = Field(description="이슈 그룹 제목")
    news_numbers: list[int] = Field(description="이 그룹에 속한 뉴스 순번")
    summary: str = Field(description="그룹 뉴스 요약. 1000자 이내")
    insights: list[str] = Field(description="주요 시사점 2~4개")

    @field_validator("summary")
    @classmethod
    def clip_summary(cls, value: str) -> str:
        text = (value or "").strip()
        return text[:1000]


class AnalyzeOutput(BaseModel):
    groups: list[NewsGroup] = Field(description="이슈별로 묶은 뉴스 그룹")
    overall_insight: str = Field(description="전체 뉴스를 관통하는 시사점 2~3문장")


class NewsState(TypedDict, total=False):
    keyword: str
    count: int
    source: str
    workdir: str
    news_items: list[dict]
    groups: list[dict]
    overall_insight: str
    keywords: list[list]
    wordcloud_path: str
    chart_path: str
    report_path: str


def get_llm() -> ChatOpenAI:
    return ChatOpenAI(model=LLM_MODEL, temperature=0.2)


def fallback_groups(items: list[dict], keyword: str) -> AnalyzeOutput:
    lines = []
    for it in items:
        lines.append(f"{it['순번']}. {it['제목']} ({it.get('날짜') or '-'}) {it.get('주요내용', '')}")
    summary = "\n".join(lines)[:1000]
    return AnalyzeOutput(
        groups=[
            NewsGroup(
                title=f"{keyword} 관련 뉴스 종합",
                news_numbers=[int(it["순번"]) for it in items],
                summary=summary,
                insights=["수집된 뉴스를 이슈별로 나누지 못해 전체 목록 기준으로 정리했습니다."],
            )
        ],
        overall_insight=f"'{keyword}' 키워드로 {len(items)}건의 뉴스를 수집했습니다.",
    )


def llm_analyze(items: list[dict], keyword: str, lg: JobLog) -> AnalyzeOutput:
    require_openai()
    payload = [
        {
            "순번": it["순번"],
            "날짜": it.get("날짜", ""),
            "제목": it["제목"],
            "주요내용": it.get("주요내용", ""),
        }
        for it in items
    ]
    news_json = json.dumps(payload, ensure_ascii=False, indent=2)
    prompt = (
        "당신은 산업 뉴스 분석가입니다. 아래 수집 뉴스를 주제 유사도 기준으로 그룹핑하세요.\n"
        "규칙:\n"
        "1. 비슷한 이슈끼리만 묶고, 그룹마다 제목을 붙이세요.\n"
        "2. summary 는 해당 그룹 뉴스만 근거로 1000자 이내로 작성합니다. 없는 사실을 만들지 마세요.\n"
        "3. insights 는 의사결정에 쓸 시사점 2~4개입니다.\n"
        "4. news_numbers 는 입력 순번만 사용합니다. 모든 뉴스가 한 그룹에는 들어가게 하세요.\n"
        f"5. 검색 키워드: {keyword}\n\n"
        f"[수집 뉴스 JSON]\n{news_json}"
    )
    lg.log("[analyze] LLM 으로 이슈 그룹핑·요약·시사점 도출")
    result = get_llm().with_structured_output(AnalyzeOutput).invoke(prompt)
    if not result.groups:
        return fallback_groups(items, keyword)
    used = {n for g in result.groups for n in g.news_numbers}
    missing = [int(it["순번"]) for it in items if int(it["순번"]) not in used]
    if missing:
        result.groups.append(
            NewsGroup(
                title="기타 뉴스",
                news_numbers=missing,
                summary="위 그룹에 넣지 못한 뉴스입니다.",
                insights=[],
            )
        )
    for g in result.groups:
        g.summary = g.summary[:1000]
    return result


# ════════════════════════════════════════════════════════════════════════════
# 6. Word 보고서 (한글 폰트)
# ════════════════════════════════════════════════════════════════════════════
def _set_rfonts(rpr, name: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if qn(attr) in rfonts.attrib:
            del rfonts.attrib[qn(attr)]


def apply_korean_fonts(doc: Document, name: str = DOCX_FONT) -> None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None and rpr_default.find(qn("w:rPr")) is not None:
            _set_rfonts(rpr_default.find(qn("w:rPr")), name)
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            style.font.name = name
            _set_rfonts(style.element.get_or_add_rPr(), name)


def shade_cell(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc: Document, header: list[str], rows: list[list[str]], widths_cm: list[float] | None = None, size: int = 9) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(cell, "12323C")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(size)
    if widths_cm:
        for r in table.rows:
            for i, w in enumerate(widths_cm):
                r.cells[i].width = Cm(w)
    doc.add_paragraph()


def add_text_block(doc: Document, text: str) -> None:
    for line in [ln.strip() for ln in str(text).splitlines() if ln.strip()]:
        if line[0] in "-•·*":
            doc.add_paragraph(line.lstrip("-•·* ").strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)


def by_number(items: list[dict]) -> dict[int, dict]:
    return {int(it["순번"]): it for it in items}


def build_word_report(state: NewsState, lg: JobLog) -> str:
    items = state["news_items"]
    groups = [NewsGroup.model_validate(g) for g in state["groups"]]
    lookup = by_number(items)
    workdir = Path(state["workdir"])

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    apply_korean_fonts(doc)
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x12, 0x32, 0x3C)
    title = f"{state['keyword']} 뉴스 분석 보고서"
    doc.core_properties.title = title
    doc.core_properties.author = "뉴스 수집 분석 에이전트"

    lg.log("[report] 표지·개요 작성")
    doc.add_heading(title, level=0)
    meta = doc.add_paragraph(
        f"작성일 {datetime.now():%Y-%m-%d %H:%M}  |  출처 {state['source']}  |  수집 {len(items)}건"
    )
    if meta.runs:
        meta.runs[0].font.color.rgb = RGBColor(0x5A, 0x6A, 0x70)

    doc.add_heading("1. 분석 개요", level=1)
    add_table(
        doc,
        ["구분", "내용"],
        [
            ["검색 키워드", state["keyword"]],
            ["뉴스 출처", state["source"]],
            ["수집 건수", f"{len(items)}건"],
            ["이슈 그룹 수", f"{len(groups)}개"],
        ],
        [5, 11],
        size=10,
    )
    if state.get("overall_insight"):
        doc.add_heading("종합 시사점", level=2)
        add_text_block(doc, state["overall_insight"])

    lg.log("[report] 이슈 그룹 섹션 작성")
    doc.add_heading("2. 이슈별 뉴스 분석", level=1)
    for i, group in enumerate(groups, start=1):
        doc.add_heading(f"2.{i} {group.title}", level=2)
        nums = ", ".join(str(n) for n in group.news_numbers) or "-"
        doc.add_paragraph(f"포함 뉴스 순번: {nums}")
        titles = [lookup[n]["제목"] for n in group.news_numbers if n in lookup]
        if titles:
            doc.add_paragraph("관련 기사")
            for t in titles:
                doc.add_paragraph(t, style="List Bullet")
        add_text_block(doc, group.summary)
        if group.insights:
            doc.add_paragraph("시사점")
            for insight in group.insights:
                doc.add_paragraph(insight, style="List Bullet")

    lg.log("[report] 키워드 그래프·워드클라우드 삽입")
    doc.add_heading("3. 주요 키워드", level=1)
    wc = state.get("wordcloud_path")
    if wc and Path(wc).exists():
        doc.add_picture(wc, width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    chart = state.get("chart_path")
    if chart and Path(chart).exists():
        doc.add_picture(chart, width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    freq = [(row[0], int(row[1])) for row in state.get("keywords") or []]
    if freq:
        add_table(
            doc,
            ["순위", "키워드", "발생 횟수"],
            [[str(i), w, f"{c:,}"] for i, (w, c) in enumerate(freq[:20], start=1)],
            [3, 8, 4],
        )

    lg.log("[report] 수집 뉴스 목록 작성")
    doc.add_heading("4. 수집 뉴스 목록", level=1)
    rows = []
    for it in items:
        rows.append(
            [
                str(it["순번"]),
                it.get("날짜") or "-",
                it["제목"][:80],
                (it.get("주요내용") or "")[:90],
                it.get("url") or "",
            ]
        )
    add_table(doc, ["순번", "날짜", "제목", "주요내용", "URL"], rows, [1.4, 2.4, 4.4, 4.4, 3.4], size=8)

    path = workdir / f"뉴스분석보고서_{datetime.now():%Y%m%d_%H%M%S}.docx"
    doc.save(str(path))
    lg.log(f"[report] 저장: {path.name}")
    return str(path)


# ════════════════════════════════════════════════════════════════════════════
# 7. LangGraph 노드
# ════════════════════════════════════════════════════════════════════════════
def build_graph(lg: JobLog):
    def news_node(state: NewsState) -> dict:
        lg.step(8, "[news] SerpAPI 뉴스 수집 시작")
        items = collect_news(state["keyword"], int(state["count"]), state["source"], lg)
        workdir = Path(state["workdir"])
        json_path = workdir / "news.json"
        json_path.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
        lg.step(32, f"[news] JSON 저장 완료 ({json_path.name})")
        return {"news_items": items}

    def analyze_node(state: NewsState) -> dict:
        items = state["news_items"]
        workdir = Path(state["workdir"])
        lg.step(38, "[analyze] 키워드 추출")
        freq = extract_keywords(news_corpus(items)).most_common(80)
        if state["keyword"] and len(state["keyword"]) >= 2:
            freq = [(w, c) for w, c in freq if w != state["keyword"]]
            freq = freq or extract_keywords(news_corpus(items)).most_common(80)

        lg.step(48, "[analyze] 워드클라우드·빈도 그래프 생성")
        wc_path = workdir / "wordcloud.png"
        chart_path = workdir / "keyword_freq.png"
        if freq:
            build_wordcloud(freq, wc_path)
            save_keyword_png(freq, chart_path)
        else:
            lg.log("[analyze] 추출 키워드가 없어 시각화를 건너뜁니다.")

        lg.step(58, "[analyze] 뉴스 그룹핑·요약·시사점")
        try:
            analyzed = llm_analyze(items, state["keyword"], lg)
        except UserError:
            raise
        except Exception as e:  # noqa: BLE001
            lg.log(f"[경고] LLM 분석 실패, 전체 그룹으로 대체합니다: {type(e).__name__}: {e}")
            analyzed = fallback_groups(items, state["keyword"])

        lg.step(78, f"[analyze] 그룹 {len(analyzed.groups)}개 도출")
        return {
            "groups": [g.model_dump() for g in analyzed.groups],
            "overall_insight": analyzed.overall_insight,
            "keywords": [[w, c] for w, c in freq],
            "wordcloud_path": str(wc_path) if wc_path.exists() else "",
            "chart_path": str(chart_path) if chart_path.exists() else "",
        }

    def report_node(state: NewsState) -> dict:
        lg.step(82, "[report] Word 보고서 생성")
        path = build_word_report(state, lg)
        lg.step(96, "[report] 다운로드 준비 완료")
        return {"report_path": path}

    graph = StateGraph(NewsState)
    graph.add_node("news", news_node)
    graph.add_node("analyze", analyze_node)
    graph.add_node("report", report_node)
    graph.add_edge(START, "news")
    graph.add_edge("news", "analyze")
    graph.add_edge("analyze", "report")
    graph.add_edge("report", END)
    return graph.compile()


# ════════════════════════════════════════════════════════════════════════════
# 8. UI 헬퍼
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class RunResult:
    keyword: str
    source: str
    news_items: list[dict]
    groups: list[dict]
    overall_insight: str
    keywords: list[tuple[str, int]]
    wordcloud_path: str
    chart_path: str
    report_path: str
    news_json: str
    workdir: Path = field(default_factory=lambda: OUTPUT_DIR)


def news_table(items: list[dict]) -> pd.DataFrame:
    cols = ["순번", "날짜", "제목", "주요내용", "url"]
    if not items:
        return pd.DataFrame(columns=cols)
    return pd.DataFrame([{c: it.get(c, "") for c in cols} for it in items])


def groups_markdown(result: RunResult) -> str:
    lines = [f"**검색어** {result.keyword} · **출처** {result.source} · **수집** {len(result.news_items)}건", ""]
    if result.overall_insight:
        lines += [f"**종합 시사점**  {result.overall_insight}", ""]
    lookup = by_number(result.news_items)
    for i, raw in enumerate(result.groups, start=1):
        g = NewsGroup.model_validate(raw)
        lines += [f"### {i}. {g.title}", ""]
        titles = [f"- [{n}] {lookup[n]['제목']}" for n in g.news_numbers if n in lookup]
        if titles:
            lines += titles + [""]
        lines += [g.summary, ""]
        if g.insights:
            lines += ["**시사점**"] + [f"- {x}" for x in g.insights] + [""]
    return "\n".join(lines)


def empty_fig() -> go.Figure:
    fig = go.Figure()
    fig.update_layout(
        template="plotly_white",
        font=dict(family=PLOTLY_FONT, color=INK),
        height=320,
        annotations=[dict(text="분석을 실행하면 빈도 그래프가 표시됩니다.", showarrow=False, font=dict(size=14))],
        xaxis=dict(visible=False),
        yaxis=dict(visible=False),
    )
    return fig


def run_pipeline(keyword: str, count: int, source: str, lg: JobLog) -> RunResult:
    keyword = (keyword or "").strip()
    if not keyword:
        raise UserError("뉴스 키워드를 입력해 주세요.")
    try:
        count = int(count)
    except (TypeError, ValueError) as e:
        raise UserError("수집 건수는 숫자여야 합니다.") from e
    if count < 1 or count > 50:
        raise UserError("수집 건수는 1~50 사이로 입력해 주세요.")
    source = source or "구글 뉴스"
    serpapi_key()
    require_openai()

    workdir = OUTPUT_DIR / f"{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}"
    workdir.mkdir(parents=True, exist_ok=True)
    lg.step(3, f"작업 폴더: {workdir.name}")
    lg.log(f"LangGraph 실행: news → analyze → report  |  {source}  |  '{keyword}'  {count}건")

    graph = build_graph(lg)
    state = graph.invoke(
        {
            "keyword": keyword,
            "count": count,
            "source": source,
            "workdir": str(workdir),
        }
    )
    items = state.get("news_items") or []
    freq = [(row[0], int(row[1])) for row in state.get("keywords") or []]
    return RunResult(
        keyword=keyword,
        source=source,
        news_items=items,
        groups=state.get("groups") or [],
        overall_insight=state.get("overall_insight") or "",
        keywords=freq,
        wordcloud_path=state.get("wordcloud_path") or "",
        chart_path=state.get("chart_path") or "",
        report_path=state.get("report_path") or "",
        news_json=json.dumps(items, ensure_ascii=False, indent=2),
        workdir=workdir,
    )


# ════════════════════════════════════════════════════════════════════════════
# 9. Gradio UI
# ════════════════════════════════════════════════════════════════════════════
CSS = f"""
.gradio-container {{
  font-family: 'Malgun Gothic','맑은 고딕','Apple SD Gothic Neo',sans-serif !important;
  max-width: 1140px !important;
}}
.news-hero {{
  background: linear-gradient(160deg, {INK} 0%, #1A4A52 62%, {TEAL} 160%);
  color: {PAPER};
  padding: 28px 30px 22px;
  border-radius: 2px;
  margin-bottom: 16px;
  border-bottom: 5px solid {GOLD};
}}
.news-hero h1 {{
  margin: 0 0 8px 0;
  font-size: 1.82rem;
  font-weight: 700;
  letter-spacing: -0.03em;
  color: {PAPER} !important;
}}
.news-hero p {{
  margin: 0;
  font-size: 0.98rem;
  line-height: 1.55;
  max-width: 46rem;
  color: #E4ECEB !important;
}}
.news-flow {{
  margin-top: 14px;
  font-size: 0.86rem;
  color: #E8D48B;
  letter-spacing: 0.02em;
}}
#log-box textarea {{
  font-family: Consolas,'Malgun Gothic',monospace;
  font-size: 12.5px;
}}
"""

HERO = """
<div class="news-hero">
  <h1 style="color:#F2F5F3;">뉴스 수집 분석 데스크</h1>
  <p style="color:#E4ECEB;">키워드와 건수만 정하면 구글·네이버 최신 뉴스를 모으고, 이슈별로 묶은 뒤 워드 보고서로 내려받습니다.</p>
  <div class="news-flow" style="color:#E8D48B;">news 수집 → analyze 그룹핑·워드클라우드 → report 워드 파일</div>
</div>
"""


def on_run(keyword, count, source):
    def work(lg: JobLog):
        return run_pipeline(keyword, count, source, lg)

    def final(res: RunResult):
        fig = keyword_bar(res.keywords) if res.keywords else empty_fig()
        wc = res.wordcloud_path if res.wordcloud_path and Path(res.wordcloud_path).exists() else None
        report = res.report_path if res.report_path and Path(res.report_path).exists() else None
        return [
            news_table(res.news_items),
            res.news_json,
            groups_markdown(res),
            wc,
            fig,
            report,
            groups_markdown(res),
        ]

    yield from stream_job(work, final, 7)


def build_ui() -> gr.Blocks:
    with gr.Blocks(title="뉴스 수집 분석 데스크") as demo:
        gr.HTML(HERO)
        with gr.Row():
            keyword_in = gr.Textbox(
                label="뉴스 키워드",
                placeholder="예: 인공지능, 반도체, 금리",
                scale=3,
            )
            count_in = gr.Slider(5, 50, value=10, step=1, label="수집 건수", scale=2)
            source_in = gr.Radio(
                ["구글 뉴스", "네이버 뉴스"],
                value="구글 뉴스",
                label="뉴스 출처",
                scale=2,
            )
        run_btn = gr.Button("수집·분석·보고서 실행", variant="primary")

        with gr.Group():
            prog = gr.HTML(progress_html(0, "대기 중"))
            log_box = gr.Textbox(
                label="실행 로그",
                lines=10,
                max_lines=10,
                interactive=False,
                autoscroll=True,
                elem_id="log-box",
            )

        with gr.Tabs():
            with gr.Tab("수집 뉴스"):
                news_df = gr.Dataframe(
                    label="순번 · 날짜 · 제목 · 주요내용 · URL",
                    interactive=False,
                    wrap=True,
                    max_height=420,
                )
                news_json = gr.Code(label="news 노드 JSON", language="json", lines=16)

            with gr.Tab("분석"):
                analyze_md = gr.Markdown()
                with gr.Row():
                    wc_img = gr.Image(label="워드클라우드", interactive=False, type="filepath", scale=1)
                    kw_plot = gr.Plot(label="주요 단어 빈도", value=empty_fig(), scale=1)

            with gr.Tab("보고서"):
                gr.Markdown("보고서 파일을 클릭하면 원하는 폴더에 저장할 수 있습니다.")
                rep_file = gr.File(label="보고서 다운로드 (.docx)", interactive=False)
                rep_md = gr.Markdown()

        common = [prog, log_box]
        run_btn.click(
            on_run,
            [keyword_in, count_in, source_in],
            common + [news_df, news_json, analyze_md, wc_img, kw_plot, rep_file, rep_md],
        )
    return demo


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(exist_ok=True)
    print(
        f"한글 폰트: {FONT_PATH} | LLM: {LLM_MODEL} | "
        f"OPENAI: {'설정됨' if os.getenv('OPENAI_API_KEY') else '없음'} | "
        f"SERPAPI: {'설정됨' if (os.getenv('SERPAPI_API_KEY') or os.getenv('SERPAPI_KEY')) else '없음'}"
    )
    build_ui().launch(
        inbrowser=True,
        theme=gr.themes.Soft(primary_hue="teal", secondary_hue="stone"),
        css=CSS,
        allowed_paths=[str(OUTPUT_DIR)],
    )
