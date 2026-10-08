"""보고서를 Markdown, PDF, Word로 만든다."""

from __future__ import annotations

import io
import re
from pathlib import Path

KOREAN_FONTS = [
    Path(r"C:\Windows\Fonts\malgun.ttf"),
    Path(r"C:\Windows\Fonts\malgunbd.ttf"),
    Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
]


def _font_path() -> Path | None:
    for path in KOREAN_FONTS:
        if path.exists():
            return path
    return None


def _blocks(markdown_text: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw in markdown_text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            blocks.append(("blank", ""))
            continue
        if line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("- "):
            blocks.append(("li", line[2:].strip()))
        else:
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            blocks.append(("p", text))
    return blocks


def to_docx_bytes(markdown_text: str) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    def _run_font(run, size: int, bold: bool = False) -> None:
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(size)
        run.bold = bold
        run.font.color.rgb = RGBColor(0x34, 0x3A, 0x40)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    for kind, text in _blocks(markdown_text):
        if kind == "blank":
            continue
        if kind == "h1":
            p = doc.add_paragraph()
            run = p.add_run(text)
            _run_font(run, 18, True)
        elif kind == "h2":
            p = doc.add_paragraph()
            run = p.add_run(text)
            _run_font(run, 14, True)
        elif kind == "h3":
            p = doc.add_paragraph()
            run = p.add_run(text)
            _run_font(run, 12, True)
        elif kind == "li":
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(text)
            _run_font(run, 11)
        else:
            p = doc.add_paragraph()
            run = p.add_run(text)
            _run_font(run, 11)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def to_pdf_bytes(markdown_text: str) -> bytes:
    from fpdf import FPDF

    font_path = _font_path()
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_left_margin(16)
    pdf.set_right_margin(16)
    if font_path:
        pdf.add_font("Malgun", fname=str(font_path))
        font = "Malgun"
    else:
        font = "Helvetica"

    def write_line(text: str, size: int) -> None:
        pdf.set_font(font, size=size)
        pdf.multi_cell(
            pdf.epw,
            max(size * 0.65, 7),
            text,
            align="L",
            new_x="LMARGIN",
            new_y="NEXT",
            wrapmode="CHAR",
        )

    for kind, text in _blocks(markdown_text):
        if kind == "blank":
            pdf.ln(3)
            continue
        if kind == "h1":
            write_line(text, 16)
            pdf.ln(2)
        elif kind == "h2":
            pdf.ln(2)
            write_line(text, 13)
        elif kind == "h3":
            write_line(text, 12)
        elif kind == "li":
            write_line(f"- {text}", 11)
        else:
            write_line(text, 11)

    return bytes(pdf.output())
